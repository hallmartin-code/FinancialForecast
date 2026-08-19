"""PPTX ingestion.

PowerPoint is the richer of the two formats: shapes carry real semantics, so titles
are identified rather than guessed, tables arrive as tables, and speaker notes are
available. Notes matter -- founders routinely leave the real numbers there.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from pitchdeck_cfo.ingest.base import Block, DeckDocument, TableBlock, TextBlock, TextKind


def _is_title(shape: Any) -> bool:
    """True for a real title placeholder, not a shape that merely looks like one."""
    try:
        return bool(shape.is_placeholder and shape.placeholder_format.idx == 0)
    except (AttributeError, ValueError):
        return False


def _flatten(shapes: Any) -> list[Any]:
    """Expand group shapes so nothing nested is silently dropped.

    Sorted by (top, left) rather than left in document order: python-pptx yields
    shapes in z-order, which on a busy slide bears no relation to reading order.
    Shapes with no position (rare, but possible) sort last rather than crashing.
    """
    out: list[Any] = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            out.extend(_flatten(shape.shapes))
        else:
            out.append(shape)
    return sorted(out, key=lambda s: (s.top if s.top is not None else 1 << 30, s.left or 0))


def _table_rows(shape: Any) -> tuple[tuple[str, ...], ...]:
    return tuple(tuple(cell.text.strip() for cell in row.cells) for row in shape.table.rows)


def _notes_text(slide: Any) -> str:
    if not slide.has_notes_slide:
        return ""
    frame = slide.notes_slide.notes_text_frame
    return frame.text.strip() if frame is not None else ""


def load_pptx(path: Path) -> DeckDocument:
    """Read a .pptx into a `DeckDocument`."""
    prs = Presentation(str(path))
    blocks: list[Block] = []
    index = 0

    for page, slide in enumerate(prs.slides, start=1):
        for shape in _flatten(slide.shapes):
            if getattr(shape, "has_table", False):
                rows = _table_rows(shape)
                if any(any(c for c in row) for row in rows):
                    blocks.append(TableBlock(index=index, page=page, rows=rows))
                    index += 1
                continue

            if not getattr(shape, "has_text_frame", False):
                continue

            # One block per paragraph: a deck's bullets are separate assertions, and
            # collapsing them into a single blob loses that.
            title_seen = False
            for para in shape.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs).strip()
                if not text:
                    continue
                kind: TextKind = "title" if (_is_title(shape) and not title_seen) else "body"
                title_seen = title_seen or kind == "title"
                blocks.append(TextBlock(index=index, page=page, kind=kind, text=text))
                index += 1

        notes = _notes_text(slide)
        if notes:
            blocks.append(TextBlock(index=index, page=page, kind="notes", text=notes))
            index += 1

    return DeckDocument(
        path=path,
        kind="pptx",
        page_count=len(prs.slides),
        blocks=tuple(blocks),
    )
