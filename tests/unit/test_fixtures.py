"""Guards on the fixture decks themselves.

These are real decks, not synthetic ones, so it is worth asserting they still have
the properties the rest of the suite relies on. A fixture that quietly changes
character would make downstream tests pass for the wrong reason.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from pypdf import PdfReader

TEXT_DECKS = ["saas_meridian.pptx", "lifesci_helion.pptx", "hybrid_nimbus.pptx"]


@pytest.mark.parametrize("name", [*TEXT_DECKS, "scanned.pdf"])
def test_fixture_exists(name: str) -> None:
    path = Path(__file__).parents[1] / "fixtures" / "decks" / name
    assert path.exists(), f"missing fixture deck: {name}"
    assert path.stat().st_size > 0


def test_scanned_fixture_yields_no_extractable_text(scanned_deck: Path) -> None:
    """The whole point of this fixture: it must look empty to a text extractor."""
    chars = sum(len(page.extract_text() or "") for page in PdfReader(scanned_deck).pages)
    assert chars == 0


def test_saas_fixture_states_the_metrics_the_suite_depends_on(saas_deck: Path) -> None:
    from pptx import Presentation

    text = " ".join(
        shape.text_frame.text
        for slide in Presentation(saas_deck).slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    for marker in ["ARR", "retention", "Gross margin", "churn"]:
        assert marker.lower() in text.lower(), f"SaaS fixture no longer states {marker}"


def test_lifesci_fixture_remains_financially_sparse(lifesci_deck: Path) -> None:
    """It earns its place by having almost no financial data. Assert that stays true."""
    from pptx import Presentation

    text = " ".join(
        shape.text_frame.text
        for slide in Presentation(lifesci_deck).slides
        for shape in slide.shapes
        if shape.has_text_frame
    ).lower()
    for absent in ["arr", "gross margin", "cac", "revenue retention"]:
        assert absent not in text, f"life-sciences fixture unexpectedly states {absent}"
