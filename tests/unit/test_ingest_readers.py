"""The two readers and the loader that dispatches between them."""

from __future__ import annotations

from pathlib import Path

import pytest

from pitchdeck_cfo.errors import (
    DeckNotFoundError,
    InsufficientTextError,
    UnsupportedFileTypeError,
)
from pitchdeck_cfo.ingest import load_deck
from pitchdeck_cfo.ingest.base import TableBlock
from pitchdeck_cfo.ingest.pdf import _is_real_table, _paragraphs


class TestPptxReader:
    def test_reads_every_slide(self, saas_deck: Path) -> None:
        d = load_deck(saas_deck)
        assert d.kind == "pptx"
        assert d.page_count == 12
        assert {b.page for b in d.blocks} == set(range(1, 13))

    def test_blocks_are_globally_ordered_without_gaps(self, saas_deck: Path) -> None:
        d = load_deck(saas_deck)
        assert [b.index for b in d.blocks] == list(range(len(d.blocks)))

    def test_pages_are_non_decreasing(self, saas_deck: Path) -> None:
        # Reading order must not jump backwards between slides.
        pages = [b.page for b in load_deck(saas_deck).blocks]
        assert pages == sorted(pages)

    def test_speaker_notes_are_preserved(self, saas_deck: Path) -> None:
        # The whole reason the pptx reader exists rather than converting to PDF.
        notes = load_deck(saas_deck).notes
        assert len(notes) == 12
        assert all(n.kind == "notes" for n in notes)

    def test_notes_come_last_on_their_slide(self, saas_deck: Path) -> None:
        d = load_deck(saas_deck)
        for note in d.notes:
            same_page = d.blocks_on(note.page)
            assert same_page[-1].index == note.index

    def test_bullets_are_separate_blocks_not_one_blob(self, saas_deck: Path) -> None:
        # Each bullet is a separate assertion the extraction pass may cite.
        d = load_deck(saas_deck)
        assert len(d.blocks) > d.page_count * 2

    def test_tables_survive_as_tables(self, hybrid_deck: Path) -> None:
        tables = load_deck(hybrid_deck).tables
        assert tables, "the hybrid fixture states its market sizes in a table"
        assert all(isinstance(t, TableBlock) for t in tables)

    def test_title_placeholders_are_detected_when_present(self, tmp_path: Path) -> None:
        # The fixtures are built from plain text boxes, so this exercises the
        # placeholder path directly rather than leaving it untested.
        from pptx import Presentation

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Acme Corp"  # type: ignore[union-attr]
        slide.placeholders[1].text_frame.text = "We sell things"  # type: ignore[union-attr]
        path = tmp_path / "titled.pptx"
        prs.save(str(path))

        kinds = {b.text: b.kind for b in load_deck(path).text_blocks}
        assert kinds["Acme Corp"] == "title"
        assert kinds["We sell things"] == "body"

    def test_ocr_flag_on_pptx_is_ignored_not_fatal(self, saas_deck: Path) -> None:
        # PPTX text is already text. Failing the run over a redundant flag is worse.
        assert load_deck(saas_deck, ocr=True).char_count > 0


class TestPdfParagraphSplitting:
    def test_blank_lines_separate_blocks(self) -> None:
        assert _paragraphs("a\nb\n\nc") == ["a\nb", "c"]

    def test_single_newlines_are_kept_inside_a_block(self) -> None:
        assert _paragraphs("wrapped\nsentence") == ["wrapped\nsentence"]

    def test_whitespace_only_lines_count_as_blank(self) -> None:
        # pypdf's layout mode pads with spaces rather than emitting empty lines.
        assert _paragraphs("a\n   \nb") == ["a", "b"]

    def test_trailing_content_is_not_dropped(self) -> None:
        assert _paragraphs("a\n\nb") == ["a", "b"]

    def test_empty_page_yields_nothing(self) -> None:
        assert _paragraphs("\n  \n") == []


class TestPdfTableFilter:
    """pdfplumber reports ruled layout as tables. Only dense ones are real."""

    def test_dense_financial_table_is_kept(self) -> None:
        rows = (
            ("$000s", "2026", "2027"),
            ("Revenue", "-", "2,430"),
            ("COGS", "-", "470"),
        )
        assert _is_real_table(rows)

    def test_sparse_layout_artifact_is_rejected(self) -> None:
        # Measured at 0.38 density on a real deck; genuine tables score ~0.8.
        rows = (("", "", ""), ("Str", "", ""), ("", "Partnership", ""))
        assert not _is_real_table(rows)

    def test_single_row_is_rejected(self) -> None:
        assert not _is_real_table((("EXIT STRATEGY", "M&A Year 3"),))

    def test_single_column_is_rejected(self) -> None:
        assert not _is_real_table((("a",), ("b",)))

    def test_empty_is_rejected_without_dividing_by_zero(self) -> None:
        assert not _is_real_table(())


class TestPdfReader:
    def test_scanned_pdf_raises_rather_than_returning_an_empty_document(
        self, scanned_deck: Path
    ) -> None:
        # Handing an empty document to the extraction pass invites the model to
        # invent a company from nothing. Failing loudly is the point.
        with pytest.raises(InsufficientTextError) as exc:
            load_deck(scanned_deck)
        assert "--ocr" in exc.value.remedy

    def test_threshold_is_configurable(self, scanned_deck: Path) -> None:
        with pytest.raises(InsufficientTextError) as exc:
            load_deck(scanned_deck, min_text_chars=1)
        assert "at least 1" in exc.value.message


class TestLoaderDispatch:
    def test_missing_file(self, tmp_path: Path) -> None:
        with pytest.raises(DeckNotFoundError):
            load_deck(tmp_path / "nope.pdf")

    def test_unsupported_suffix(self, tmp_path: Path) -> None:
        p = tmp_path / "notes.txt"
        p.write_text("hello")
        with pytest.raises(UnsupportedFileTypeError):
            load_deck(p)

    def test_suffix_matching_is_case_insensitive(self, tmp_path: Path, saas_deck: Path) -> None:
        upper = tmp_path / "DECK.PPTX"
        upper.write_bytes(saas_deck.read_bytes())
        assert load_deck(upper).kind == "pptx"


class TestOCRGating:
    def test_missing_binaries_produce_install_instructions(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        from pitchdeck_cfo.errors import OCRUnavailableError
        from pitchdeck_cfo.ingest import ocr

        monkeypatch.setattr(ocr.shutil, "which", lambda _: None)
        with pytest.raises(OCRUnavailableError) as exc:
            ocr.check_available()
        assert "tesseract" in exc.value.message
        assert "poppler" in exc.value.message
        assert "PATH" in exc.value.remedy
