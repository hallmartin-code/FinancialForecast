"""The workbook. Structure, formula-not-value, and -- the one that matters -- that
the Excel formulas evaluate to the same numbers as the Python model.

The model is implemented twice, once in `model/` and once in Excel formulas, and the
whole value of the workbook rests on the two agreeing. `TestExcelAgreesWithPython`
compiles the workbook with a spreadsheet engine and checks every P&L line and the
cash balance. Those tests are marked `slow` because compiling takes a few seconds.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from pitchdeck_cfo.model import build
from pitchdeck_cfo.render import workbook
from tests import model_factories as mf

ENGINES = ["saas", "hardware", "life_sciences"]

# Deliberately non-trivial, so no assertion can pass because everything is zero.
FULL = dict(
    cogs=mf.cogs(
        hosting_pct_of_revenue=mf.s(9.0),
        hosting_scale_exponent=mf.s(0.85),
        support_pct_of_revenue=mf.s(4.0),
        payment_processing_pct=mf.s(2.9),
        bom_pct_of_asp=mf.s(40.0),
        consumable_gross_margin_pct=mf.s(75.0),
        warranty_pct_of_revenue=mf.s(2.0),
    ),
    headcount=mf.headcount(
        starting={"Engineering": 6, "Sales": 2, "G&A": 1},
        loaded_multiplier=mf.s(1.28),
        annual_attrition_pct=mf.s(15.0),
        rep_quota_annual=mf.s(600_000),
        rep_ramp_months=mf.s(6.0),
        quota_attainment_pct=mf.s(75.0),
        accounts_per_csm=mf.s(40.0),
        engineers_per_product_line=mf.s(6.0),
        product_lines=mf.s(1.0),
        revenue_per_engineer=mf.s(400_000),
        sales_per_marketing_hire=mf.s(3.0),
        ftes_per_ga_hire=mf.s(12.0),
        scale_exponent=mf.s(0.75),
    ),
    opex=mf.opex(
        marketing_pct_of_new_revenue=mf.s(40.0),
        rd_tooling_per_engineer=mf.s(6_000),
        rent_per_fte=mf.s(6_000),
        software_per_fte=mf.s(3_600),
        ga_fixed_annual=mf.s(180_000),
    ),
    working_capital=mf.working_capital(
        dso_days=mf.s(45.0),
        dpo_days=mf.s(30.0),
        inventory_days=mf.s(20.0),
        capex_pct_of_revenue=mf.s(2.0),
        capex_depreciation_years=mf.s(3.0),
    ),
    financing=mf.financing(cash_on_hand=mf.s(1_000_000), raise_amount=mf.s(6_000_000)),
    tax=mf.tax(blended_rate_pct=mf.s(25.0)),
)


def _model(engine: str) -> Any:
    revenue = {
        "saas": lambda: mf.saas_revenue(
            new_logo_growth_monthly_pct=mf.s(4.0),
            growth_decay_annual_pct=mf.s(30.0),
            terminal_growth_monthly_pct=mf.s(1.0),
            logo_churn_annual_pct=mf.s(10.0),
            net_revenue_retention_pct=mf.s(112.0),
            annual_prepay_mix_pct=mf.s(50.0),
        ),
        "hardware": lambda: mf.hardware_revenue(
            unit_growth_monthly_pct=mf.s(5.0),
            growth_decay_annual_pct=mf.s(30.0),
            terminal_growth_monthly_pct=mf.s(1.0),
            direct_sales_mix_pct=mf.s(60.0),
            distributor_margin_pct=mf.s(25.0),
        ),
        "life_sciences": mf.life_sciences_revenue,
    }[engine]
    return build(mf.assumptions(engine, revenue=revenue(), **FULL))  # type: ignore[arg-type]


@pytest.fixture(scope="module")
def saas_book(tmp_path_factory: pytest.TempPathFactory) -> Path:
    out = tmp_path_factory.mktemp("wb") / "saas.xlsx"
    return workbook.write(_model("saas"), out)


class TestStructure:
    def test_every_expected_tab_is_present_and_ordered(self, saas_book: Path) -> None:
        assert load_workbook(saas_book).sheetnames == list(workbook.SHEETS)

    def test_assumptions_become_defined_names(self, saas_book: Path) -> None:
        # Formulas reference assumptions by name, which is what makes them readable.
        names = set(load_workbook(saas_book).defined_names)
        assert "revenue_arpu_annual" in names
        assert "tax_blended_rate_pct" in names
        assert len(names) > 40

    def test_month_and_year_headers_exist(self, saas_book: Path) -> None:
        ws = load_workbook(saas_book)["PnL"]
        assert ws.cell(row=3, column=3).value.startswith("20")
        assert ws.freeze_panes is not None

    def test_the_readme_states_provenance_and_the_disclaimer(self, saas_book: Path) -> None:
        text = "\n".join(str(c.value) for c in load_workbook(saas_book)["README"]["A"] if c.value)
        assert "core inputs sourced from the deck" in text
        assert "not audited financial information" in text
        assert "Blue cells" in text


class TestFormulasNotValues:
    """A workbook of pasted numbers cannot be stress-tested, so it does not do the job."""

    @pytest.mark.parametrize(
        ("sheet", "label"),
        [
            ("PnL", "Revenue"),
            ("PnL", "Gross Profit"),
            ("PnL", "EBITDA"),
            ("PnL", "Net Income"),
            ("Cashflow", "Ending Cash"),
            ("COGS", "Total COGS"),
            ("Opex", "Total operating expense"),
            ("Headcount", "Total headcount"),
        ],
    )
    def test_key_lines_are_live_formulas(self, saas_book: Path, sheet: str, label: str) -> None:
        ws = load_workbook(saas_book)[sheet]
        row = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value == label)
        for column in (3, 20, 62):
            value = ws.cell(row=row, column=column).value
            assert isinstance(value, str) and value.startswith("="), (
                f"{sheet}!{label} column {column} is a literal, not a formula"
            )

    def test_assumption_values_are_literals(self, saas_book: Path) -> None:
        # The inputs are the one place a number belongs. Everything else derives.
        ws = load_workbook(saas_book)["Assumptions"]
        row = next(
            r
            for r in range(5, ws.max_row + 1)
            if ws.cell(row=r, column=1).value == "tax.blended_rate_pct"
        )
        assert ws.cell(row=row, column=2).value == 25.0

    def test_inputs_are_blue_and_formulas_are_not(self, saas_book: Path) -> None:
        # The convention a financial analyst reads without being told.
        wb = load_workbook(saas_book)
        ws = wb["Assumptions"]
        row = next(
            r
            for r in range(5, ws.max_row + 1)
            if ws.cell(row=r, column=1).value == "tax.blended_rate_pct"
        )
        assert "0000CC" in str(ws.cell(row=row, column=2).font.color.rgb)

        # A formula cell is anything but blue -- often no explicit colour at all.
        pnl = wb["PnL"]
        colour = pnl.cell(row=4, column=3).font.color
        assert colour is None or "0000CC" not in str(colour.rgb)

    def test_no_formula_references_a_missing_name(self, saas_book: Path) -> None:
        wb = load_workbook(saas_book)
        known = set(wb.defined_names)
        import re

        pattern = re.compile(r"\b([a-z][a-z0-9_]{4,})\b")
        functions = {"if", "max", "min", "sum", "ceiling", "average", "sumproduct", "offset"}
        for sheet in ("Revenue", "Headcount", "COGS", "Opex", "PnL", "Cashflow", "Metrics"):
            ws = wb[sheet]
            for row in ws.iter_rows():
                for cell in row:
                    if not isinstance(cell.value, str) or not cell.value.startswith("="):
                        continue
                    for token in pattern.findall(cell.value):
                        if token in functions or token in known:
                            continue
                        raise AssertionError(
                            f"{sheet}!{cell.coordinate} references unknown name {token!r}"
                        )


# --------------------------------------------------------------------------- #
# the test that actually matters
# --------------------------------------------------------------------------- #

formulas = pytest.importorskip("formulas", reason="spreadsheet engine not installed")


def _evaluate(path: Path) -> Any:
    import logging
    import warnings

    warnings.filterwarnings("ignore")
    logging.disable(logging.WARNING)
    solution = formulas.ExcelModel().loads(str(path.resolve())).finish().calculate()

    def value(sheet: str, cell: str) -> Any:
        import numpy as np

        raw = np.asarray(solution[f"'[{path.name}]{sheet}'!{cell}"].value).ravel()[0]
        return raw if isinstance(raw, str) else float(raw)

    return value


def _annual_rows(path: Path, sheet: str = "PnL") -> dict[str, int]:
    ws = load_workbook(path)[sheet]
    rows: dict[str, int] = {}
    seen = False
    for r in range(1, ws.max_row + 1):
        label = ws.cell(row=r, column=1).value
        if label == "Annual roll-up":
            seen = True
            continue
        if seen and isinstance(label, str):
            rows.setdefault(label, r)
    return rows


PNL_LINES = [
    "Revenue",
    "COGS",
    "Gross Profit",
    "R&D",
    "S&M",
    "G&A",
    "Total Opex",
    "EBITDA",
    "Depreciation",
    "Tax",
    "Net Income",
]


@pytest.mark.slow
class TestExcelAgreesWithPython:
    """The model is implemented twice. If the two disagree, the workbook is a lie."""

    @pytest.mark.parametrize("engine", ENGINES)
    def test_every_pnl_line_matches(self, engine: str, tmp_path: Path) -> None:
        model = _model(engine)
        path = workbook.write(model, tmp_path / f"{engine}.xlsx")
        value = _evaluate(path)
        rows = _annual_rows(path)

        for line in PNL_LINES:
            for i, column in enumerate("CDEFG"):
                excel = value("PNL", f"{column}{rows[line]}")
                assert not isinstance(excel, str), f"{line} Y{i + 1} evaluated to {excel!r}"
                assert excel == pytest.approx(model.pnl_annual[line][i], abs=0.01), (
                    f"{engine} {line} Y{i + 1}"
                )

    @pytest.mark.parametrize("engine", ENGINES)
    def test_ending_cash_matches(self, engine: str, tmp_path: Path) -> None:
        model = _model(engine)
        path = workbook.write(model, tmp_path / f"{engine}.xlsx")
        value = _evaluate(path)
        ws = load_workbook(path)["Cashflow"]
        row = [
            r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value == "Ending Cash"
        ][-1]
        for i, column in enumerate("CDEFG"):
            assert value("CASHFLOW", f"{column}{row}") == pytest.approx(
                model.ending_cash[i], abs=0.01
            )

    def test_changing_an_assumption_recalculates_the_model(self, tmp_path: Path) -> None:
        """The whole point of the workbook: it is a live model, not a screenshot."""
        model = _model("saas")
        path = workbook.write(model, tmp_path / "before.xlsx")
        rows = _annual_rows(path)
        before = _evaluate(path)("PNL", f"G{rows['Revenue']}")

        wb = load_workbook(path)
        ws = wb["Assumptions"]
        target = next(
            r
            for r in range(5, ws.max_row + 1)
            if ws.cell(row=r, column=1).value == "revenue.net_revenue_retention_pct"
        )
        ws.cell(row=target, column=2, value=130.0)
        edited = tmp_path / "after.xlsx"
        wb.save(edited)

        after = _evaluate(edited)("PNL", f"G{rows['Revenue']}")
        assert after > before * 1.1, "raising NRR did not move Y5 revenue"
